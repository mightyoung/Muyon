"""Permit-checked python-pptx copy source; trusted worker bootstrap required."""
from .runner.worker_gate import require_worker_permit


def edit_copy(job,source_path,candidate_path,*,permit=None):
    admitted=require_worker_permit(permit,job,source_path,candidate_path,'python-pptx')
    job=admitted.request
    from pptx import Presentation
    admitted.engine_entered()
    document=Presentation(source_path);op=job['operation']
    slides=[s for s in document.slides if s.slide_id == op['slide_id']]
    if len(slides)!=1:raise ValueError('slide_identity_conflict')
    shapes=[s for s in slides[0].shapes if s.shape_id == op['shape_id']]
    if len(shapes)!=1:raise ValueError('shape_identity_conflict')
    run=shapes[0].text_frame.paragraphs[0].runs[0]
    if run.text != op['expected']:raise ValueError('expected_value_conflict')
    run.text=op['replacement']
    # Saving may lose/normalize unsupported parts; strict independent oracle
    # rejects them. No actual engine round-trip has been performed.
    with candidate_path.open('xb') as output:document.save(output)


def main():
    from .runner.worker_entry import dispatch
    return dispatch('pptx')

if __name__=='__main__':
    import json
    print(json.dumps(main(),sort_keys=True))
